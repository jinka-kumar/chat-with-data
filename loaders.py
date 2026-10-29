"""
File loaders: turn PDF, Word, Excel, CSV and text files into "sections".

Every loader returns the same shape, so the rest of the app doesn't care about file types:
    sections: [(section_no, text), ...]       -> fed to chunk_pages() as before
    labels:   {section_no: "Sheet1, rows 2-21"} -> human-friendly citation for each section

Why sections and not pages? Only PDFs have real pages. Word has headings,
Excel has sheets and rows, text files have lines. The label says where to look.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROWS_PER_SECTION = 20       # spreadsheet rows grouped per citation ("rows 2-21")
LINES_PER_SECTION = 40      # text-file lines grouped per citation ("lines 1-40")
PARAS_PER_SECTION = 25      # Word paragraphs per section when there are no headings
MAX_ROWS = 5_000            # protect the laptop: embedding 100k rows locally takes hours

SUPPORTED = ["pdf", "docx", "xlsx", "csv", "txt", "md"]

Sections = list[tuple[int, str]]
Labels = dict[int, str]


class UnsupportedFile(ValueError):
    """Raised for file types we can't read (e.g. old .doc / .xls)."""


def read_file(name: str, data: bytes) -> tuple[Sections, Labels]:
    """Pick the right loader from the file extension."""
    ext = Path(name).suffix.lower().lstrip(".")
    loaders = {"pdf": _read_pdf, "docx": _read_docx, "xlsx": _read_xlsx,
               "csv": _read_csv, "txt": _read_text, "md": _read_text}
    if ext not in loaders:
        hint = " Save it as .docx / .xlsx first." if ext in ("doc", "xls") else ""
        raise UnsupportedFile(f"'.{ext}' files are not supported.{hint}")
    return loaders[ext](data)


# --- PDF: real pages ---------------------------------------------------------
def _read_pdf(data: bytes) -> tuple[Sections, Labels]:
    sections = []
    for page_no, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append((page_no, text))
    return sections, {n: f"p. {n}" for n, _ in sections}


# --- Word: split by headings (fallback: every N paragraphs) -----------------
def _read_docx(data: bytes) -> tuple[Sections, Labels]:
    doc = Document(io.BytesIO(data))
    sections: Sections = []
    labels: Labels = {}
    title, parts, count = "Start", [], 0

    def flush():
        nonlocal parts, count
        text = "\n\n".join(p for p in parts if p.strip())
        if text:
            n = len(sections) + 1
            sections.append((n, text))
            labels[n] = f"section '{title}'"
        parts, count = [], 0

    # iter_inner_content() keeps paragraphs AND tables in document order.
    for block in doc.iter_inner_content():
        if hasattr(block, "rows"):                       # a table
            parts.extend(_table_rows_as_text([[c.text for c in r.cells] for r in block.rows]))
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.startswith("Heading") or style == "Title":
            flush()
            title = block.text.strip()[:60] or title
            parts.append(block.text)
        else:
            parts.append(block.text)
            count += 1
            if count >= PARAS_PER_SECTION:               # long doc without headings
                flush()
    flush()
    return sections, labels


# --- Excel / CSV: rows as "column: value" text -------------------------------
def _table_rows_as_text(rows: list[list]) -> list[str]:
    """First row = header. Each data row becomes one self-describing line:
    'Region: North; Product: Laptop; Units: 5'. A blank line between rows
    (added by the caller) keeps the chunker from cutting a row in half."""
    if not rows:
        return []
    header = [str(h).strip() if h not in (None, "") else f"Column {i + 1}" for i, h in enumerate(rows[0])]
    lines = []
    for row in rows[1:]:
        cells = [f"{h}: {v}" for h, v in zip(header, row) if v not in (None, "")]
        if cells:
            lines.append("; ".join(cells) + ".")
    return lines


def _rows_to_sections(sheet: str, rows: list[list], sections: Sections, labels: Labels) -> None:
    lines = _table_rows_as_text(rows)
    for i in range(0, len(lines), ROWS_PER_SECTION):
        block = lines[i : i + ROWS_PER_SECTION]
        n = len(sections) + 1
        sections.append((n, "\n\n".join(block)))
        first = i + 2                                    # +2: header is row 1 in Excel
        labels[n] = f"{sheet}, rows {first}-{first + len(block) - 1}"


def _read_xlsx(data: bytes) -> tuple[Sections, Labels]:
    # read_only + data_only: fast, low memory, and formulas give their computed values.
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: Sections = []
    labels: Labels = {}
    for ws in wb.worksheets:
        rows = [list(r) for _, r in zip(range(MAX_ROWS + 1), ws.iter_rows(values_only=True))]
        _rows_to_sections(ws.title, rows, sections, labels)
    wb.close()
    return sections, labels


def _read_csv(data: bytes) -> tuple[Sections, Labels]:
    text = data.decode("utf-8-sig", errors="replace")    # utf-8-sig drops Excel's BOM
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel                              # default: comma-separated
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for _, r in zip(range(MAX_ROWS + 1), reader)]
    sections: Sections = []
    labels: Labels = {}
    _rows_to_sections("CSV", rows, sections, labels)
    return sections, labels


# --- Plain text / Markdown: groups of lines ----------------------------------
def _read_text(data: bytes) -> tuple[Sections, Labels]:
    lines = data.decode("utf-8", errors="replace").splitlines()
    sections: Sections = []
    labels: Labels = {}
    for i in range(0, len(lines), LINES_PER_SECTION):
        text = "\n".join(lines[i : i + LINES_PER_SECTION]).strip()
        if text:
            n = len(sections) + 1
            sections.append((n, text))
            labels[n] = f"lines {i + 1}-{min(i + LINES_PER_SECTION, len(lines))}"
    return sections, labels
