"""
Data agent: answers number questions over spreadsheets with SQL (DuckDB).

Why SQL instead of RAG for tables?
RAG shows the LLM ~4 chunks (a keyhole view), so it cannot total 5,000 rows.
Here the LLM only WRITES a query; DuckDB RUNS it on all rows, exactly.

Flow:  pick relevant columns -> LLM writes SQL -> safety check -> run read-only
       -> (on error, LLM fixes it once) -> LLM explains the result.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import tempfile
import threading
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterator

import duckdb
import ollama
from openpyxl import load_workbook

from planner import PLAN_PROMPT, PlanError, build_sql, parse_plan, validate_plan
from rag import CHAT_MODEL, DB_PATH, bm25_scores, cite, tokenize

# A code-focused model writes better SQL, e.g. SQL_MODEL=qwen2.5-coder:7b (optional).
SQL_MODEL = os.getenv("SQL_MODEL", "qwen2.5-coder:14b")
TABLE_EXTENSIONS = ("csv", "xlsx")
MAX_COLUMNS_IN_PROMPT = 25     # wide sheets: only the most relevant columns go to the LLM
SAMPLE_ROWS = 3                # a few example rows help the LLM see value formats
MAX_RESULT_ROWS = 50           # results are for people to read, not to dump the table
QUERY_TIMEOUT_S = 15           # stop runaway queries (e.g. accidental huge cross joins)


# --- Prompts -------------------------------------------------------------------
SQL_PROMPT = """You write ONE DuckDB SQL query that answers the user's question.

Tables you can use. Each column has a note saying what it contains: its type, its exact
values or examples, and its range. READ THESE NOTES FIRST.
{schema}

Rules:
- Use ONE table whenever it has all the columns you need. JOIN only if a needed column exists
  only in another table. Never invent columns: use EXACTLY the names listed for each table.
- Match text values exactly as listed in the notes, e.g. "category" = 'Phones'.
- For "most / top / highest / lowest" questions, return the name AND the number used to rank it,
  e.g. SELECT "category", SUM("units") AS total_units ... ORDER BY total_units DESC LIMIT 1.
- Return ONLY the SQL. No explanation, no markdown.
- A single SELECT (or WITH ... SELECT) statement. Never modify data.
- Wrap column and table names in double quotes, e.g. "units".
- Text matching should be case-insensitive: use ILIKE, e.g. "region" ILIKE 'north'.
- Return only the columns needed to answer. Add ORDER BY for "top"/"highest" questions.
- Date columns are TIMESTAMP. Filter dates like "order_date" BETWEEN '2026-07-01' AND '2026-09-30 23:59:59'.
- Never invent values. Only use numbers from the question or the facts below."""

FACTS_PROMPT = """

Facts from the user's documents (use them for definitions such as date ranges or thresholds):
{facts}

If the question compares data with a target or limit from these facts:
- Do NOT put the target numbers in the query and do NOT use subqueries.
- Return the ACTUAL total per group. The comparison is done in the next step. Use this shape:
  SELECT "<group column>", SUM("<measure column>") AS actual FROM "<table>"
  WHERE "<date column>" BETWEEN '<start date>' AND '<end date> 23:59:59' GROUP BY 1 ORDER BY 1"""

PLAN_FACTS_PROMPT = """

Facts from the user's documents (use them for date ranges and definitions):
{facts}
If the question compares data with a target or limit from these facts, plan the ACTUAL metric
per group (e.g. sum of units per region in that date range). Never put target numbers in the plan."""

EMPTY_PROMPT = """Your query returned NO rows, so it cannot be right.
Return the actual total per group, with no target numbers and no subqueries, like:
SELECT "<group column>", SUM("<measure column>") AS actual FROM "<table>" WHERE <date filter> GROUP BY 1 ORDER BY 1
Return ONLY the SQL."""

EXTRACT_TARGETS_PROMPT = """From the document text, find the target number for each of these groups: {groups}.
Reply with JSON only, like {{"North": 190, "South": null}}. Use null if the text has no target for a group.

Document text:
{docs}"""

FIX_PROMPT = """Your SQL failed. Write a corrected query.
SQL: {sql}
Error: {error}
Return ONLY the corrected SQL."""

EXPLAIN_PROMPT = """You explain the result of a SQL query to a business user.
Rules:
- Use ONLY the values in the result below. Never invent, round wildly or estimate numbers.
- Answer the question in 1-4 short sentences or bullets. Format big numbers with commas.
- If the result is empty, say that no matching rows were found.
- End with the source in brackets, like [{sources}, SQL]."""

COMBINED_PROMPT = """Answer using BOTH sources below and nothing else.
- Facts from documents: cite them like [policy.pdf, p. 3].
- Numbers from the data: use ONLY the SQL result and cite like [sales.xlsx, SQL]. Never invent numbers.
- If a COMPARISON table is given, it was calculated exactly by code. Report every row as it is
  (met or missed). Never change, add or drop a group.
- If the data result has no rows or failed, say that the actual numbers could not be found.
  In that case NEVER say which groups met or missed a target.

Document context:
{docs}

Data result (SQL: {sql}):
{data}

COMPARISON (actual vs target, calculated by code):
{comparison}"""


# --- Data classes ----------------------------------------------------------------
@dataclass
class Table:
    name: str                 # SQL name the LLM uses, e.g. "sales"
    source: str               # file name, e.g. "sales.xlsx"
    sheet: str                # sheet name ("CSV" for csv files)
    path: str                 # parquet file on disk
    rows: int
    columns: list[str]
    types: list[str]
    sample: list[dict] = field(default_factory=list)


@dataclass
class DataResult:
    question: str
    sql: str = ""
    columns: list[str] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    truncated: bool = False
    sources: list[str] = field(default_factory=list)
    error: str | None = None
    plan: str = ""            # plain-English plan, e.g. "number of different salesperson"

    @property
    def ok(self) -> bool:
        return self.error is None


class UnsafeSQL(ValueError):
    """The generated SQL is not a single read-only SELECT."""


# --- SQL safety ----------------------------------------------------------------
# Defence in depth: (1) this text check, (2) the DuckDB connection itself is locked
# so it cannot read/write files even if something slips through.
_FORBIDDEN = re.compile(
    r"\b(insert|update|delete|drop|create|alter|attach|detach|copy|export|import|install|"
    r"load|pragma|set|reset|call|checkpoint|vacuum|grant|revoke|truncate|use)\b", re.I)
_FILE_FUNCS = re.compile(r"\b(read_\w+|glob|sniff_csv|parquet_\w+|query_table)\s*\(", re.I)


def validate_sql(sql: str) -> str:
    """Return a cleaned SQL string, or raise UnsafeSQL."""
    sql = sql.strip().rstrip(";").strip()
    # Hide string literals and quoted names so 'Update Ltd' or "update_date" don't trigger.
    masked = re.sub(r"'(?:[^']|'')*'", "''", sql)
    masked = re.sub(r'"(?:[^"]|"")*"', '""', masked)
    if not sql:
        raise UnsafeSQL("Empty query")
    if ";" in masked:
        raise UnsafeSQL("Only one statement is allowed")
    if not re.match(r"^\s*(select|with)\b", masked, re.I):
        raise UnsafeSQL("Only SELECT queries are allowed")
    if m := _FORBIDDEN.search(masked):
        raise UnsafeSQL(f"Keyword not allowed: {m.group(0).upper()}")
    if m := _FILE_FUNCS.search(masked):
        raise UnsafeSQL(f"File access not allowed: {m.group(1)}")
    return sql


def extract_sql(text: str) -> str:
    """Small models wrap SQL in ```sql fences or add chatter. Pull out just the query."""
    if m := re.search(r"```(?:sql)?\s*(.*?)```", text, re.S | re.I):
        text = m.group(1)
    if m := re.search(r"\b(with|select)\b.*", text, re.S | re.I):
        text = m.group(0)
    return text.split(";")[0].strip()


# --- Column selection for wide tables ---------------------------------------------
def select_columns(question: str, table: Table, k: int = MAX_COLUMNS_IN_PROMPT) -> list[str]:
    """1000 columns don't fit in a prompt. Keep the first column (usually an ID/name)
    plus the columns whose NAMES best match the question (BM25 on column names)."""
    cols = table.columns
    if len(cols) <= k:
        return cols
    docs = [tokenize(c) + tokenize(c.replace("_", " ")) for c in cols]
    scores = bm25_scores(tokenize(question.replace("_", " ")) + tokenize(question), docs)
    ranked = sorted(range(1, len(cols)), key=lambda i: (-scores[i], i))   # stable: ties keep order
    keep = sorted([0] + ranked[: k - 1])
    return [cols[i] for i in keep]


def schema_text(question: str, tables: list[Table], notes: dict | None = None) -> str:
    """The 'data brief' the SQL writer reads first: every column with its type and a note
    (exact values, examples or range) so the model knows what each column really holds."""
    notes = notes or {}
    parts = []
    for t in tables:
        cols = select_columns(question, t)
        lines = [f'  "{c}" {ty}' + (f"  -- {notes[(t.name, c)]}" if (t.name, c) in notes else "")
                 for c, ty in zip(t.columns, t.types) if c in cols]
        hidden = len(t.columns) - len(cols)
        if hidden:
            lines.append(f"  (+{hidden} more columns not shown)")
        sample = [{c: _short(r.get(c)) for c in cols} for r in t.sample[:2]]
        parts.append(f'TABLE "{t.name}"  -- from {t.source}, sheet {t.sheet}, {t.rows:,} rows\n'
                     + "\n".join(lines) + f"\n  example rows: {json.dumps(sample, default=str)}")
    return "\n\n".join(parts)


def _words(text: str) -> set[str]:
    """Lowercase words, split on _ too, with a simple plural strip: 'Phones' -> 'phone'."""
    return {w[:-1] if len(w) > 3 and w.endswith("s") else w for w in re.findall(r"[a-z0-9]+", str(text).lower())}


def select_tables(question: str, tables: list[Table], notes: dict) -> list[Table]:
    """Send only the sheets that match the question: words in the table name, column names
    or column values ('phones' is a value of "category"). Fewer tables = fewer bad JOINs."""
    if len(tables) <= 1:
        return tables
    q = _words(question)
    scores = []
    for t in tables:
        vocab = _words(t.name.replace(_safe_name(Path(t.source).stem), "")) | _words(" ".join(t.columns))
        vocab |= _words(" ".join(v for (name, _), v in notes.items() if name == t.name))
        scores.append(len(q & vocab))
    best = max(scores)
    if best == 0:
        return tables
    return [t for t, s in zip(tables, scores) if s >= 0.6 * best]


def _short(v, n: int = 30):
    s = str(v) if v is not None else None
    return s if s is None or len(s) <= n else s[:n] + "..."


def _safe_name(text: str) -> str:
    name = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_") or "table"
    return f"t_{name}" if name[0].isdigit() else name


# --- Storage: spreadsheets -> parquet + a small JSON catalog ------------------------
class DataStore:
    """Each uploaded sheet is saved once as Parquet (fast, typed, compressed).
    Queries load the selected sheets into an in-memory, locked DuckDB connection."""

    def __init__(self, path: str | Path = Path(DB_PATH) / "tables"):
        self.dir = Path(path)
        self.dir.mkdir(parents=True, exist_ok=True)
        self._con_key: tuple | None = None
        self._con: duckdb.DuckDBPyConnection | None = None
        self._notes: dict[tuple[str, str], str] = {}   # (parquet path, column) -> profile note

    def _catalog(self, doc_id: str) -> Path:
        return self.dir / f"{doc_id}.json"

    def has(self, doc_id: str) -> bool:
        return self._catalog(doc_id).exists()

    def tables(self, doc_id: str) -> list[Table]:
        return [Table(**t) for t in json.loads(self._catalog(doc_id).read_text(encoding="utf-8"))]

    def list_files(self) -> dict[str, str]:
        """{doc_id: file name} for every stored spreadsheet."""
        out = {}
        for p in sorted(self.dir.glob("*.json")):
            tables = json.loads(p.read_text(encoding="utf-8"))
            if tables:
                out[p.stem] = tables[0]["source"]
        return out

    def add(self, doc_id: str, source: str, data: bytes) -> list[Table]:
        ext = Path(source).suffix.lower().lstrip(".")
        if ext not in TABLE_EXTENSIONS:
            raise ValueError(f"Not a spreadsheet: {source}")
        sheets = _csv_sheets(data) if ext == "csv" else _xlsx_sheets(data)
        base = _safe_name(Path(source).stem)
        tables: list[Table] = []
        con = duckdb.connect()
        with tempfile.TemporaryDirectory() as tmp:
            for i, (sheet, csv_text) in enumerate(sheets):
                src = Path(tmp) / f"{i}.csv"
                src.write_text(csv_text, encoding="utf-8")
                dst = self.dir / f"{doc_id}__{i}.parquet"
                read = f"read_csv({_lit(src)}, header=true, sample_size=-1)"   # -1: infer types from ALL rows
                # Our own clean names ("Unit Price ($)" -> unit_price). DuckDB's normalize_names
                # renames reserved words like "Name" to "_name", which confuses the LLM.
                originals = [d[0] for d in con.execute(f"DESCRIBE SELECT * FROM {read}").fetchall()]
                select = ", ".join(f'"{_q(o)}" AS "{c}"' for o, c in zip(originals, clean_columns(originals)))
                con.execute(f"COPY (SELECT {select} FROM {read}) TO {_lit(dst)} (FORMAT parquet)")
                desc = con.execute(f"DESCRIBE SELECT * FROM read_parquet({_lit(dst)})").fetchall()
                cur = con.execute(f"SELECT * FROM read_parquet({_lit(dst)}) LIMIT {SAMPLE_ROWS}")
                names = [d[0] for d in cur.description]
                sample = [dict(zip(names, r)) for r in cur.fetchall()]
                rows = con.execute(f"SELECT count(*) FROM read_parquet({_lit(dst)})").fetchone()[0]
                name = base if ext == "csv" or len(sheets) == 1 else f"{base}_{_safe_name(sheet)}"
                tables.append(Table(name=name, source=source, sheet=sheet, path=str(dst), rows=rows,
                                    columns=[d[0] for d in desc], types=[d[1] for d in desc],
                                    sample=json.loads(json.dumps(sample, default=str))))
        con.close()
        if not tables:
            raise ValueError("No data found (empty file?)")
        self._catalog(doc_id).write_text(json.dumps([asdict(t) for t in tables]), encoding="utf-8")
        return tables

    # -- querying --
    def connection(self, doc_ids: list[str]) -> tuple[duckdb.DuckDBPyConnection, list[Table]]:
        """Locked, in-memory DuckDB with the selected sheets loaded as tables.
        Cached for the last selection so follow-up questions are instant."""
        tables = [t for d in doc_ids if self.has(d) for t in self.tables(d)]
        seen: dict[str, int] = {}
        for t in tables:                       # two files with the same name: sales, sales_2
            seen[t.name] = seen.get(t.name, 0) + 1
            if seen[t.name] > 1:
                t.name = f"{t.name}_{seen[t.name]}"
        key = tuple(sorted(doc_ids))
        if key != self._con_key:
            if self._con is not None:
                self._con.close()
            con = duckdb.connect(":memory:")
            for t in tables:
                con.execute(f'CREATE TABLE "{t.name}" AS SELECT * FROM read_parquet({_lit(t.path)})')
            # Lock it down: no file or network access, and settings can't be changed back.
            con.execute("SET enable_external_access = false")
            con.execute("SET lock_configuration = true")
            self._con, self._con_key = con, key
        return self._con, tables

    def column_notes(self, tables: list[Table], question: str) -> dict[tuple[str, str], str]:
        """Profile the columns the SQL writer will see (cached). Pure SQL, no LLM:
        text -> exact values (or examples), numbers -> min/max, dates -> first/last."""
        con = self._con
        out = {}
        for t in tables:
            for c, ty in zip(t.columns, t.types):
                if c not in select_columns(question, t):
                    continue
                key = (t.path, c)
                if key not in self._notes:
                    self._notes[key] = _profile_column(con, t.name, c, ty)
                out[(t.name, c)] = self._notes[key]
        return out

    def run_sql(self, sql: str, doc_ids: list[str]) -> tuple[list[str], list[list], bool]:
        sql = validate_sql(sql)
        con, _ = self.connection(doc_ids)
        cur = con.cursor()
        timer = threading.Timer(QUERY_TIMEOUT_S, cur.interrupt)   # stop runaway queries
        timer.start()
        try:
            cur.execute(f"SELECT * FROM ({sql}) AS q LIMIT {MAX_RESULT_ROWS + 1}")
            columns = [d[0] for d in cur.description]
            rows = [list(r) for r in cur.fetchall()]
        finally:
            timer.cancel()
            cur.close()
        truncated = len(rows) > MAX_RESULT_ROWS
        return columns, rows[:MAX_RESULT_ROWS], truncated

    def ask(self, question: str, doc_ids: list[str], facts: list[str] | None = None) -> DataResult:
        """Question -> SQL -> result. One automatic fix attempt if the SQL fails.
        facts: document text for the "both" route (targets, date ranges, limits)."""
        _, tables = self.connection(doc_ids)
        result = DataResult(question=question, sources=sorted({t.source for t in tables}))
        if not tables:
            result.error = "No spreadsheets selected."
            return result
        # Understand the data first: profile the columns, then keep only the matching sheets.
        notes = self.column_notes(tables, question)
        tables = select_tables(question, tables, notes)
        schema = schema_text(question, tables, notes)
        facts_text = "\n".join(f"- {f[:400]}" for f in (facts or [])[:4])

        # Step 1: understand the question as a PLAN (a small JSON form); code writes the SQL.
        try:
            plan_prompt = PLAN_PROMPT.format(schema=schema)
            if facts:
                plan_prompt += PLAN_FACTS_PROMPT.format(facts=facts_text)
            plan, table = validate_plan(parse_plan(_llm(plan_prompt, question, model=SQL_MODEL)), tables)
            result.sql = build_sql(plan, table)
            result.columns, result.rows, result.truncated = self.run_sql(result.sql, doc_ids)
            if result.rows or not facts:
                result.plan = plan.describe()
                return result
        except (PlanError, UnsafeSQL, duckdb.Error, ValueError, TypeError, KeyError, AttributeError):
            pass                       # plan didn't fit: fall back to free-form SQL below
        except Exception as exc:       # model not reachable
            result.error = f"Could not reach the model: {exc}"
            return result
        result.columns, result.rows, result.sql = [], [], ""

        # Step 2 (fallback): free-form SQL for complex questions (joins, subqueries, ...).
        system = SQL_PROMPT.format(schema=schema)
        if facts:
            system += FACTS_PROMPT.format(facts=facts_text)
        try:
            sql = extract_sql(_llm(system, question, model=SQL_MODEL))
        except Exception as exc:
            result.error = f"Could not reach the model: {exc}"
            return result
        for attempt in range(2):
            result.sql = sql
            try:
                result.columns, result.rows, result.truncated = self.run_sql(sql, doc_ids)
                result.error = None
                if result.rows or not facts or attempt == 1:
                    return result
                # "both" questions should always return actual values. Empty = wrong query: retry once.
                sql = extract_sql(_llm(system, EMPTY_PROMPT, model=SQL_MODEL))
                continue
            except (UnsafeSQL, duckdb.Error) as exc:
                result.error = f"{type(exc).__name__}: {exc}"
                if attempt == 0:   # let the model fix its own mistake once
                    try:
                        sql = extract_sql(_llm(system, FIX_PROMPT.format(sql=sql, error=exc), model=SQL_MODEL))
                    except Exception:
                        break
        return result


def _profile_column(con, table: str, col: str, ty: str, max_values: int = 15) -> str:
    q, tbl = f'"{_q(col)}"', f'"{_q(table)}"'
    try:
        if ty.startswith(("TIMESTAMP", "DATE")):
            lo, hi = con.execute(f"SELECT min({q})::DATE, max({q})::DATE FROM {tbl}").fetchone()
            return f"date, from {lo} to {hi}"
        if any(k in ty for k in ("INT", "DOUBLE", "FLOAT", "DECIMAL", "REAL", "NUMERIC")):
            lo, hi = con.execute(f"SELECT min({q}), max({q}) FROM {tbl}").fetchone()
            return f"number, min {lo}, max {hi}"
        n = con.execute(f"SELECT count(DISTINCT {q}) FROM {tbl}").fetchone()[0]
        vals = [str(r[0]) for r in con.execute(
            f"SELECT {q} FROM {tbl} WHERE {q} IS NOT NULL GROUP BY 1 ORDER BY count(*) DESC, 1 LIMIT {max_values}").fetchall()]
        vals = [_short(v, 40) for v in vals]
        if n <= max_values:
            return f"text, {n} values: " + ", ".join(vals)
        return f"text, {n:,} different values, e.g. " + ", ".join(vals[:5])
    except duckdb.Error:
        return ty.lower()


def clean_columns(names: list[str]) -> list[str]:
    """Spreadsheet headers -> simple, unique snake_case names: 'Unit Price ($)' -> 'unit_price',
    '2026 Sales' -> 'c_2026_sales', blank -> 'column_3', duplicates -> 'region_2'."""
    out: list[str] = []
    for i, n in enumerate(names, start=1):
        c = re.sub(r"[^a-z0-9]+", "_", str(n).lower()).strip("_") or f"column_{i}"
        if c[0].isdigit():
            c = f"c_{c}"
        base, k = c, 2
        while c in out:
            c, k = f"{base}_{k}", k + 1
        out.append(c)
    return out


def _q(name: str) -> str:
    """Escape a double quote inside a quoted SQL identifier."""
    return str(name).replace('"', '""')


def _lit(path: str | Path) -> str:
    """Quote a file path as a SQL string literal."""
    return "'" + str(path).replace("'", "''") + "'"


def _csv_sheets(data: bytes) -> list[tuple[str, str]]:
    return [("CSV", data.decode("utf-8-sig", errors="replace"))]


def _xlsx_sheets(data: bytes) -> list[tuple[str, str]]:
    """Each non-empty sheet -> CSV text. data_only=True gives formula RESULTS, not formulas."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        buf = io.StringIO()
        writer = csv.writer(buf)
        count = 0
        for row in ws.iter_rows(values_only=True):
            if any(v not in (None, "") for v in row):
                writer.writerow(["" if v is None else v for v in row])
                count += 1
        if count >= 2:                          # need a header + at least one row
            sheets.append((ws.title, buf.getvalue()))
    wb.close()
    return sheets


# --- LLM helpers -----------------------------------------------------------------
def _llm(system: str, user: str, model: str = CHAT_MODEL) -> str:
    resp = ollama.chat(model=model, options={"temperature": 0},
                       messages=[{"role": "system", "content": system}, {"role": "user", "content": user}])
    return resp["message"]["content"].strip()


def _stream(system: str, user: str) -> Iterator[str]:
    for part in ollama.chat(model=CHAT_MODEL, stream=True, options={"temperature": 0.1},
                            messages=[{"role": "system", "content": system}, {"role": "user", "content": user}]):
        yield part["message"]["content"]


def result_as_text(result: DataResult, max_rows: int = 20) -> str:
    """Compact table for the LLM: header + up to 20 rows."""
    if not result.ok:
        return f"(query failed: {result.error})"
    if not result.rows:
        return "(no rows)"
    lines = [" | ".join(result.columns)]
    lines += [" | ".join("" if v is None else str(v) for v in r) for r in result.rows[:max_rows]]
    if len(result.rows) > max_rows or result.truncated:
        lines.append(f"... (more rows not shown)")
    return "\n".join(lines)


# --- Number check: every number in an answer must come from a source ------------------
_NUMBER = re.compile(r"(?<![\w.])-?\d[\d,]*(?:\.\d+)?")


def _norm(n: str) -> str:
    n = n.replace(",", "")
    return n.rstrip("0").rstrip(".") if "." in n else n


def ungrounded_numbers(answer: str, sources: list[str]) -> list[str]:
    """Numbers in the answer that appear in none of the sources. Small counts (0-10, e.g.
    "3 regions") are allowed; rounding to whole numbers or 1-2 decimals is allowed."""
    allowed: set[str] = set()
    for text in sources:
        for raw in _NUMBER.findall(text):
            n = _norm(raw)
            allowed.add(n)
            try:
                x = float(n)
                allowed |= {str(round(x)), _norm(f"{x:.1f}"), _norm(f"{x:.2f}")}
            except ValueError:
                pass
    bad = []
    for raw in _NUMBER.findall(answer):
        n = _norm(raw)
        if n in allowed or (n.lstrip("-").isdigit() and abs(int(n)) <= 10):
            continue
        bad.append(raw)
    return bad


def table_answer(result: DataResult) -> str:
    """Plain answer built from the result table (used when the AI's wording isn't trustworthy)."""
    src = f"[{', '.join(result.sources) or 'data'}, SQL]"
    if not result.ok:
        return "I couldn't answer that from the data. Try naming the column or sheet."
    if not result.rows:
        return f"No matching rows were found in the data. {src}"
    if len(result.rows) == 1:
        return "Result: " + ", ".join(f"**{c}**: {v}" for c, v in zip(result.columns, result.rows[0])) + f" {src}"
    lines = [" · ".join(f"{c}: {v}" for c, v in zip(result.columns, r)) for r in result.rows[:10]]
    more = f"\n\n(+{len(result.rows) - 10} more rows in the table below)" if len(result.rows) > 10 else ""
    return "Here is what the data shows:\n\n" + "\n".join(f"- {l}" for l in lines) + more + f"\n\n{src}"


def explain_answer(result: DataResult) -> tuple[str, bool]:
    """(answer, checked). checked=False means the AI's wording used numbers that are not in
    the result, so we show a plain answer built from the table instead."""
    if not result.ok or not result.rows:
        return table_answer(result), True
    system = EXPLAIN_PROMPT.format(sources=", ".join(result.sources) or "data")
    user = f"Question: {result.question}\nSQL: {result.sql}\nResult:\n{result_as_text(result)}"
    try:
        text = _llm(system, user)
    except Exception:
        text = ""
    if text and not ungrounded_numbers(text, [result_as_text(result), result.question]):
        return text, True
    return table_answer(result), False


def explain_stream(result: DataResult) -> Iterator[str]:
    system = EXPLAIN_PROMPT.format(sources=", ".join(result.sources) or "data")
    user = f"Question: {result.question}\nSQL: {result.sql}\nResult:\n{result_as_text(result)}"
    yield from _stream(system, user)


# --- Targets: compare in CODE, not in the LLM -----------------------------------------
# Small models get "who met the target" wrong surprisingly often (and sound sure about it).
# So: numbers come from SQL, targets come from the document text, and Python compares them.
_TARGET_WORDS = re.compile(r"\b(targets?|quotas?|goals?)\b", re.I)


def _num(v) -> float | None:
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def _fmt(x: float) -> str:
    return f"{x:,.0f}" if float(x).is_integer() else f"{x:,.2f}"


def extract_targets(groups: list[str], docs_text: str) -> dict[str, float]:
    """Find '<group> ... target ... <number>' in the text (no LLM). Then, for groups still
    missing, ask the LLM, but only accept numbers that literally appear in the text."""
    targets: dict[str, float] = {}
    for g in groups:
        # The gap between the group name and "target" must NOT mention another group. Otherwise
        # "East 140 ... West 190 ... The North region target is 190" would give East = 190.
        others = "|".join(re.escape(o) for o in groups if o != g) or "(?!)"
        gap = rf"(?:(?!\b(?:{others})\b)[^.\n]){{0,80}}?"
        word, number = r"\b(?:target|quota|goal)\b", r"(\d[\d,]*(?:\.\d+)?)"
        patterns = [
            rf"\b{re.escape(g)}\b{gap}{word}[^\d.\n]{{0,30}}{number}",   # "North region ... target is 190"
            rf"{word}{gap}\b{re.escape(g)}\b[^\d.\n]{{0,30}}{number}",   # "target for North is 190"
        ]
        for pattern in patterns:
            if m := re.search(pattern, docs_text, re.I):
                targets[g] = _num(m.group(1))
                break
    missing = [g for g in groups if g not in targets]
    if missing:
        try:
            raw = _llm(EXTRACT_TARGETS_PROMPT.format(groups=", ".join(missing), docs=docs_text[:4000]), "JSON:")
            found = json.loads(re.search(r"\{.*\}", raw, re.S).group(0))
        except Exception:
            found = {}
        lower = {str(k).strip().lower(): v for k, v in found.items()}
        plain_text = docs_text.replace(",", "")
        for g in missing:
            n = _num(lower.get(g.lower()))
            # Anti-hallucination: the number must appear in the document text.
            if n is not None and re.search(rf"(?<![\d.]){re.escape(_fmt(n).replace(',', ''))}(?![\d])", plain_text):
                targets[g] = n
    return targets


def compare_with_targets(question: str, hits: list[dict], result: DataResult) -> list[dict] | None:
    """Rows of {group, actual, target, difference, status}, or None if not a target question
    or the data doesn't look like 'group | number'."""
    if not _TARGET_WORDS.search(question) or not result.ok or not result.rows or len(result.columns) < 2:
        return None
    actuals = [(str(r[0]), _num(r[-1])) for r in result.rows]
    if any(a is None for _, a in actuals):
        return None
    targets = extract_targets([g for g, _ in actuals], "\n".join(h["text"] for h in hits))
    if not targets:
        return None
    rows = []
    for g, actual in actuals:
        t = targets.get(g)
        rows.append({"group": g, "actual": _fmt(actual), "target": _fmt(t) if t is not None else "not found",
                     "difference": f"{actual - t:+,.0f}" if t is not None else "",
                     "status": ("met" if actual >= t else "missed") if t is not None else "unknown"})
    return rows


def comparison_as_text(rows: list[dict] | None) -> str:
    if not rows:
        return "(none)"
    return "\n".join(f"{r['group']}: actual {r['actual']}, target {r['target']}, "
                     f"difference {r['difference'] or '-'}, {r['status'].upper()}" for r in rows)


def combined_answer(question: str, hits: list[dict], result: DataResult,
                    comparison: list[dict] | None = None) -> tuple[str, bool]:
    """Like combined_stream, but checks every number against the documents, the data and
    the comparison. If the AI invented a number, fall back to a plain, code-built answer."""
    docs = "\n\n".join(f"[{cite(h)}] {h['text']}" for h in hits) or "(nothing found in documents)"
    system = COMBINED_PROMPT.format(docs=docs, sql=result.sql or "-", data=result_as_text(result),
                                    comparison=comparison_as_text(comparison))
    try:
        text = _llm(system, question)
    except Exception:
        text = ""
    sources = [docs, result_as_text(result), comparison_as_text(comparison), question]
    if text and not ungrounded_numbers(text, sources):
        return text, True
    if comparison:
        src = f"[{', '.join(result.sources) or 'data'}, SQL]"
        met = [r["group"] for r in comparison if r["status"] == "met"]
        missed = [r["group"] for r in comparison if r["status"] == "missed"]
        return (f"**Met the target:** {', '.join(met) or 'none'}  \n**Missed the target:** "
                f"{', '.join(missed) or 'none'}  \nSee the comparison table below. {src}"), False
    return table_answer(result), False


def combined_stream(question: str, hits: list[dict], result: DataResult,
                    comparison: list[dict] | None = None) -> Iterator[str]:
    docs = "\n\n".join(f"[{cite(h)}] {h['text']}" for h in hits) or "(nothing found in documents)"
    system = COMBINED_PROMPT.format(docs=docs, sql=result.sql or "-", data=result_as_text(result),
                                    comparison=comparison_as_text(comparison))
    yield from _stream(system, question)
