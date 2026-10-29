"""Loader tests: build real Word / Excel / CSV / text files in memory, then read them back."""
import io

import pytest
from docx import Document
from openpyxl import Workbook

import rag
from loaders import UnsupportedFile, read_file
from rag import PdfStore, chunk_pages


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Leave Policy", level=1)
    doc.add_paragraph("Staff get 20 days of annual leave.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Type", "Days"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Sick", "12"
    doc.add_heading("Travel Policy", level=1)
    doc.add_paragraph("Book flights 14 days ahead.")
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active; ws.title = "Sales"
    ws.append(["Region", "Product", "Units"])
    for i in range(25):
        ws.append(["North" if i % 2 else "South", f"SKU-{i:03d}", i])
    wb.create_sheet("Notes").append(["Note"])
    wb["Notes"].append(["Q3 target is 500 units"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def test_docx_sections_follow_headings_and_keep_tables():
    sections, labels = read_file("hr.docx", make_docx())
    assert list(labels.values()) == ["section 'Leave Policy'", "section 'Travel Policy'"]
    assert "Type: Sick; Days: 12." in sections[0][1]          # table row kept, readable
    assert "Book flights" in sections[1][1]


def test_xlsx_rows_become_text_with_excel_row_numbers():
    sections, labels = read_file("sales.xlsx", make_xlsx())
    assert labels[1] == "Sales, rows 2-21" and labels[2] == "Sales, rows 22-26"
    assert "Region: North; Product: SKU-001; Units: 1." in sections[0][1]
    assert labels[3] == "Notes, rows 2-2" and "Q3 target" in sections[2][1]


def test_csv_semicolons_and_excel_bom():
    data = "﻿name;city\nAsha;Hyderabad\nRavi;Pune\n".encode("utf-8")
    sections, labels = read_file("people.csv", data)
    assert sections[0][1] == "name: Asha; city: Hyderabad.\n\nname: Ravi; city: Pune."
    assert labels[1] == "CSV, rows 2-3"


def test_text_and_markdown_grouped_by_lines():
    data = "\n".join(f"line {i}" for i in range(1, 51)).encode()
    sections, labels = read_file("notes.md", data)
    assert labels == {1: "lines 1-40", 2: "lines 41-50"}


@pytest.mark.parametrize("name", ["old.doc", "old.xls", "image.png"])
def test_unsupported_files_raise_clear_error(name):
    with pytest.raises(UnsupportedFile):
        read_file(name, b"whatever")


def test_rows_are_never_cut_in_half():
    sections, _ = read_file("sales.xlsx", make_xlsx())
    for chunk in chunk_pages(sections, size=120, overlap=30):
        assert chunk.text.endswith(".")                       # every chunk ends at a row end


def test_end_to_end_citation_uses_sheet_and_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(rag, "embed", lambda texts, kind: [[t.count("SKU"), 1.0] for t in texts])
    sections, labels = read_file("sales.xlsx", make_xlsx())
    store = PdfStore(str(tmp_path))
    store.add("x", "sales.xlsx", chunk_pages(sections), labels=labels)
    hit = store.search("SKU-024", ["x"], k=1)[0]
    assert "SKU-024" in hit["text"]
    assert rag.cite(hit) == "sales.xlsx, Sales, rows 22-26"


def test_old_data_without_loc_still_cites_page():
    assert rag.cite({"source": "a.pdf", "page": 7}) == "a.pdf, p. 7"
