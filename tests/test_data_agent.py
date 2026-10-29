"""Data agent tests: SQL safety, storage, wide tables, and the ask() flow with a fake LLM."""
import csv
from datetime import datetime
import io

import duckdb
import pytest
from openpyxl import Workbook

import data_agent
from data_agent import (DataStore, UnsafeSQL, extract_sql, result_as_text, select_columns,
                        validate_sql)


# ---------- test files ----------
def sales_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Product", "Units", "Unit Price ($)"])
    rows = [("North", "Laptop", 5, 1000), ("South", "Laptop", 3, 1000), ("North", "Phone", 10, 500),
            ("East", "Phone", 7, 500), ("South", "Tablet", 4, 300)]
    for r in rows:
        ws.append(r)
    targets = wb.create_sheet("Targets")
    targets.append(["Region", "Target"])
    for r in [("North", 9000), ("South", 5000), ("East", 4000)]:
        targets.append(r)
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def wide_csv(n_rows=50, n_cols=1000) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["customer_id"] + [f"metric_{i:04d}" for i in range(1, n_cols)])
    for r in range(n_rows):
        w.writerow([f"CUST-{r:04d}"] + [r * 10 + i for i in range(1, n_cols)])
    return buf.getvalue().encode()


@pytest.fixture
def store(tmp_path):
    s = DataStore(tmp_path / "tables")
    s.add("x1", "sales.xlsx", sales_xlsx())
    return s


# ---------- SQL safety ----------
@pytest.mark.parametrize("sql", [
    "SELECT 1",
    "select region, sum(units) from sales group by 1",
    "WITH t AS (SELECT 1 AS a) SELECT a FROM t",
    "SELECT * FROM sales WHERE product = 'Update; Drop'",       # keywords inside strings are fine
    'SELECT "update_date", "set" FROM sales',                   # ... and inside quoted names
    "SELECT replace(product, 'a', 'b') FROM sales;",            # trailing ; is trimmed
])
def test_safe_sql_allowed(sql):
    assert validate_sql(sql)


@pytest.mark.parametrize("sql", [
    "DROP TABLE sales",
    "DELETE FROM sales",
    "SELECT 1; DROP TABLE sales",
    "UPDATE sales SET units = 0",
    "SELECT * FROM read_csv('/etc/passwd')",
    "SELECT * FROM read_text('~/.ssh/id_rsa')",
    "COPY sales TO '/tmp/leak.csv'",
    "ATTACH '/tmp/x.db' AS x",
    "INSTALL httpfs",
    "SET enable_external_access = true",
    "",
])
def test_unsafe_sql_blocked(sql):
    with pytest.raises(UnsafeSQL):
        validate_sql(sql)


def test_connection_is_locked_even_if_validator_is_bypassed(store):
    con, _ = store.connection(["x1"])
    with pytest.raises(duckdb.Error):
        con.execute("SELECT * FROM read_csv_auto('/etc/hostname')")      # file access disabled
    with pytest.raises(duckdb.Error):
        con.execute("SET enable_external_access = true")                 # settings locked


def test_extract_sql_from_chatty_model():
    text = "Sure! Here is the query:\n```sql\nSELECT 1;\n```\nThis counts rows."
    assert extract_sql(text) == "SELECT 1"
    assert extract_sql("SELECT a FROM t; -- done") == "SELECT a FROM t"


# ---------- storage ----------
def test_xlsx_sheets_become_tables_with_clean_column_names(store):
    tables = {t.sheet: t for t in store.tables("x1")}
    assert set(tables) == {"Sales", "Targets"}                           # empty sheet skipped
    sales = tables["Sales"]
    assert sales.name == "sales_sales" and sales.rows == 5
    assert sales.columns[:4] == ["region", "product", "units", "unit_price"]
    assert store.list_files() == {"x1": "sales.xlsx"}


def test_run_sql_gives_exact_numbers(store):
    cols, rows, truncated = store.run_sql(
        'SELECT "region", sum("units" * "unit_price") AS revenue FROM "sales_sales" '
        'GROUP BY 1 ORDER BY 2 DESC', ["x1"])
    assert cols == ["region", "revenue"]
    assert rows == [["North", 10000], ["South", 4200], ["East", 3500]] and not truncated


def test_results_are_capped(tmp_path):
    s = DataStore(tmp_path)
    s.add("w", "wide.csv", wide_csv(n_rows=120, n_cols=3))
    _, rows, truncated = s.run_sql('SELECT * FROM "wide"', ["w"])
    assert len(rows) == data_agent.MAX_RESULT_ROWS and truncated


def test_same_file_name_twice_gets_unique_table_names(tmp_path):
    s = DataStore(tmp_path)
    s.add("a", "sales.csv", b"x\n1\n")
    s.add("b", "sales.csv", b"x\n2\n")
    _, tables = s.connection(["a", "b"])
    assert sorted(t.name for t in tables) == ["sales", "sales_2"]


# ---------- wide tables ----------
def test_select_columns_keeps_id_and_relevant_columns(tmp_path):
    s = DataStore(tmp_path)
    t = s.add("w", "wide.csv", wide_csv())[0]
    assert len(t.columns) == 1000
    cols = select_columns("which customer has the highest metric_0500?", t, k=25)
    assert cols[0] == "customer_id" and "metric_0500" in cols and len(cols) == 25


def test_wide_table_query_end_to_end(tmp_path, monkeypatch):
    s = DataStore(tmp_path)
    s.add("w", "wide.csv", wide_csv())
    prompts = []

    def fake_llm(system, user, model=None):
        prompts.append(system)
        return 'SELECT "customer_id", "metric_0500" FROM "wide" ORDER BY "metric_0500" DESC LIMIT 1'
    monkeypatch.setattr(data_agent, "_llm", fake_llm)
    res = s.ask("Which customer has the highest metric_0500?", ["w"])
    assert res.ok and res.rows == [["CUST-0049", 990]]
    assert '"metric_0500"' in prompts[0] and "(+975 more columns not shown)" in prompts[0]
    assert len(prompts[0]) < 6000                       # prompt stays small even with 1000 columns


# ---------- ask(): retry and failure ----------
def test_ask_fixes_bad_sql_once(store, monkeypatch):
    replies = iter(['SELECT "revenue_total" FROM "sales_sales"',                     # bad column
                    'SELECT sum("units") AS units FROM "sales_sales"'])              # fixed
    monkeypatch.setattr(data_agent, "_llm", lambda system, user, model=None: next(replies))
    res = store.ask("How many units in total?", ["x1"])
    assert res.ok and res.rows == [[29]] and res.sources == ["sales.xlsx"]


def test_ask_reports_error_after_second_failure(store, monkeypatch):
    monkeypatch.setattr(data_agent, "_llm", lambda system, user, model=None: "DROP TABLE sales_sales")
    res = store.ask("delete everything", ["x1"])
    assert not res.ok and "UnsafeSQL" in res.error


def test_ask_with_model_down(store, monkeypatch):
    def down(*a, **k):
        raise ConnectionError("ollama not running")
    monkeypatch.setattr(data_agent, "_llm", down)
    res = store.ask("total units?", ["x1"])
    assert not res.ok and "Could not reach the model" in res.error


def test_result_as_text():
    r = data_agent.DataResult("q", sql="SELECT 1", columns=["a", "b"], rows=[[1, None]])
    assert result_as_text(r) == "a | b\n1 | "
    assert result_as_text(data_agent.DataResult("q", error="boom")) == "(query failed: boom)"


# ---------- column names (regression: DuckDB turned "Name" into "_name") ----------
def test_clean_columns():
    assert data_agent.clean_columns(["Unit Price ($)", "2026 Sales", "", "Region", "region", 'He said "hi"']) == \
        ["unit_price", "c_2026_sales", "column_3", "region", "region_2", "he_said_hi"]


def test_reserved_word_headers_stay_readable(tmp_path):
    s = DataStore(tmp_path)
    t = s.add("e", "staff.csv", b"Name,Role,Order,Group\nAsha,Lead,1,A\nRavi,Dev,2,B\n")[0]
    assert t.columns == ["name", "role", "order", "group"]
    _, rows, _ = s.run_sql('SELECT "name" FROM "staff" WHERE "order" = 2', ["e"])
    assert rows == [["Ravi"]]


# ---------- "both" route: document facts reach the SQL writer ----------
def test_facts_are_given_to_sql_writer(store, monkeypatch):
    prompts = []

    def fake_llm(system, user, model=None):
        prompts.append(system)
        return ('SELECT "region", sum("units") AS units FROM "sales_sales" GROUP BY 1 ORDER BY 1')
    monkeypatch.setattr(data_agent, "_llm", fake_llm)
    res = store.ask("Which regions met their Q3 target?", ["x1"],
                    facts=["The Q3 2026 period runs from 1 July 2026 to 30 September 2026.",
                           "The North region Q3 2026 target is 190 units."])
    assert res.ok
    plan_prompt, sql_prompt = prompts[0], prompts[1]   # fake reply isn't a plan -> SQL fallback
    for p in (plan_prompt, sql_prompt):
        assert "1 July 2026 to 30 September 2026" in p and "190 units" in p
    assert "plan the ACTUAL metric" in plan_prompt and "Return the ACTUAL total per group" in sql_prompt


def test_no_facts_section_for_plain_data_questions(store, monkeypatch):
    prompts = []
    monkeypatch.setattr(data_agent, "_llm", lambda s, u, model=None: prompts.append(s) or "SELECT 1")
    store.ask("How many rows?", ["x1"])
    assert "Facts from the user's documents" not in prompts[0]


# ---------- targets are compared in code, not by the LLM ----------
HANDBOOK = ("Region Q3 2026 target (units) Regional Manager North 190 Arjun Rao South 200 Divya Nair "
            "East 140 Ankit Das West 190 Rohan Patil The North region Q3 2026 target is 190 units. "
            "The South region Q3 2026 target is 200 units. The East region Q3 2026 target is 140 units. "
            "The West region Q3 2026 target is 190 units.")
GROUPS = ["North", "South", "East", "West"]


def no_llm(monkeypatch):
    def down(*a, **k):
        raise ConnectionError
    monkeypatch.setattr(data_agent, "_llm", down)


def test_extract_targets_from_real_handbook_text(monkeypatch):
    no_llm(monkeypatch)
    # Regression: East was read as 190 (North's number) when the pattern ran across groups.
    assert data_agent.extract_targets(GROUPS, HANDBOOK) == {"North": 190, "South": 200, "East": 140, "West": 190}


@pytest.mark.parametrize("text, expected", [
    ("The Q3 target for North is 20 units. The target for South is 35 units.", {"North": 20, "South": 35}),
    ("North: quarterly quota of 1,200 units. South: quarterly quota of 900 units.", {"North": 1200, "South": 900}),
])
def test_extract_targets_word_order_variants(monkeypatch, text, expected):
    no_llm(monkeypatch)
    assert data_agent.extract_targets(["North", "South"], text) == expected


def test_llm_targets_must_appear_in_text(monkeypatch):
    text = "Quarterly goals table: Alpha 1,250 | Beta 900"          # no "<group> ... target <n>" sentence
    monkeypatch.setattr(data_agent, "_llm",
                        lambda *a, **k: '{"Alpha": 1250, "Beta": 950, "Gamma": 10}')   # Beta 950 is invented
    assert data_agent.extract_targets(["Alpha", "Beta", "Gamma"], text) == {"Alpha": 1250}


def q3_result(rows):
    return data_agent.DataResult("q", sql="SELECT ...", columns=["region", "actual"], rows=rows)


def test_compare_with_targets_met_and_missed(monkeypatch):
    no_llm(monkeypatch)
    rows = data_agent.compare_with_targets(
        "Which regions met their Q3 target?", [{"text": HANDBOOK}],
        q3_result([["East", 143], ["North", 204], ["South", 180], ["West", 172]]))
    assert [(r["group"], r["status"], r["difference"]) for r in rows] == [
        ("East", "met", "+3"), ("North", "met", "+14"), ("South", "missed", "-20"), ("West", "missed", "-18")]


@pytest.mark.parametrize("question, result", [
    ("Which orders break the discount policy?", q3_result([["East", 143]])),    # not a target question
    ("Which regions met their target?", q3_result([])),                           # no data
    ("Which regions met their target?", data_agent.DataResult("q", error="boom")),
    ("Which regions met their target?", q3_result([["East", "n/a"]])),           # last column not numeric
])
def test_no_comparison_when_not_applicable(monkeypatch, question, result):
    no_llm(monkeypatch)
    assert data_agent.compare_with_targets(question, [{"text": HANDBOOK}], result) is None


def test_empty_result_retries_once_for_both_route(store, monkeypatch):
    replies = iter(['SELECT "region" FROM "sales_sales" WHERE "units" > 9999',          # returns nothing
                    'SELECT "region", sum("units") AS actual FROM "sales_sales" GROUP BY 1 ORDER BY 1'])
    monkeypatch.setattr(data_agent, "_llm", lambda s, u, model=None: next(replies))
    res = store.ask("Which regions met their target?", ["x1"], facts=["North target is 9 units."])
    assert res.rows == [["East", 7], ["North", 15], ["South", 7]]


def test_combined_prompt_forbids_claims_without_data(monkeypatch):
    seen = {}

    def fake_stream(system, user):
        seen["system"] = system
        yield "ok"
    monkeypatch.setattr(data_agent, "_stream", fake_stream)
    list(data_agent.combined_stream("Which regions met?", [], q3_result([]), None))
    assert "NEVER say which groups met or missed" in seen["system"] and "(no rows)" in seen["system"]


# ---------- understand the data first: column profiles + table selection ----------
def bluepeak_like() -> bytes:
    wb = Workbook()
    s = wb.active
    s.title = "Sales"
    s.append(["Order Date", "Region", "Salesperson", "Category", "Units", "Unit Price"])
    s.append([datetime(2026, 1, 5), "North", "Asha", "Phones", 2, 100])
    s.append([datetime(2026, 9, 30), "South", "Ravi", "Laptops", 1, 900])
    s.append([datetime(2026, 3, 1), "North", "Ravi", "Phones", 4, 100])
    e = wb.create_sheet("Employees")
    e.append(["Name", "Role", "Leave Taken Days"])
    e.append(["Asha", "Exec", 20])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def blue(tmp_path):
    s = DataStore(tmp_path / "t")
    s.add("b", "shop.xlsx", bluepeak_like())
    return s


def test_column_notes_describe_values_ranges_and_dates(blue):
    _, tables = blue.connection(["b"])
    notes = blue.column_notes(tables, "anything")
    assert notes[("shop_sales", "category")] == "text, 2 values: Phones, Laptops"
    assert notes[("shop_sales", "units")] == "number, min 1, max 4"
    assert notes[("shop_sales", "order_date")] == "date, from 2026-01-05 to 2026-09-30"


def test_only_matching_sheets_are_sent(blue, monkeypatch):
    prompts = []
    monkeypatch.setattr(data_agent, "_llm", lambda s, u, model=None: prompts.append(s) or "SELECT 1")
    blue.ask("which sales person did more sales on phones", ["b"])
    for p in prompts:            # plan prompt AND the free-SQL fallback both get the data brief
        assert '"shop_sales"' in p and '"shop_employees"' not in p   # no tempting JOIN
        assert "Phones, Laptops" in p


# ---------- number check: no invented numbers ----------
@pytest.mark.parametrize("answer, bad", [
    ("Phones sold the most: 123,456 units.", ["123,456"]),        # the real hallucination we saw
    ("Phones sold the most: 568 units.", []),
    ("Kabir Shah sold 63,59,250 (Indian format).", []),
    ("Top 3 categories; 2 of them are devices.", []),            # small counts are fine
    ("Average was 1,234.57.", []),                               # rounding allowed
])
def test_ungrounded_numbers(answer, bad):
    sources = ["category | total_units\nPhones | 568", "salesperson | total\nKabir Shah | 6359250.0",
               "avg\n1234.5678"]
    assert data_agent.ungrounded_numbers(answer, sources) == bad


def test_invented_number_falls_back_to_table_answer(monkeypatch):
    monkeypatch.setattr(data_agent, "_llm", lambda s, u, model=None: "Phones sold 123,456 units.")
    r = data_agent.DataResult("q", sql="...", columns=["category", "total_units"], rows=[["Phones", 568]],
                              sources=["s.xlsx"])
    text, checked = data_agent.explain_answer(r)
    assert not checked and "568" in text and "123,456" not in text


def test_combined_answer_falls_back_to_comparison(monkeypatch):
    monkeypatch.setattr(data_agent, "_llm", lambda s, u, model=None: "North and West met (West sold 999).")
    r = data_agent.DataResult("q", sql="...", columns=["region", "actual"], rows=[["North", 204], ["West", 172]])
    comp = [{"group": "North", "actual": "204", "target": "190", "difference": "+14", "status": "met"},
            {"group": "West", "actual": "172", "target": "190", "difference": "-18", "status": "missed"}]
    text, checked = data_agent.combined_answer("Which met?", [{"text": "targets", "page": 2, "source": "h.pdf"}], r, comp)
    assert not checked and "**Met the target:** North" in text and "**Missed the target:** West" in text
